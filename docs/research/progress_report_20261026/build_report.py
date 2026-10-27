from pathlib import Path
import json, re, hashlib, datetime
import pandas as pd
import bibtexparser
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START

O=Path(__file__).parent
R=Path('/root/workspace/Code/DualTimesField')
D=R/'logs/forecasting/main_comparison_20261023'
baselines=json.loads((O/'baselines.json').read_text())
method=json.loads((O/'method.json').read_text())
audit=json.loads((O/'results_audit.json').read_text())
main=pd.read_csv(D/'main.csv'); reg=pd.read_csv(D/'regional.csv')
refs=[]; refmap={}
def addref(key,citation,url=''):
 if key not in refmap:refmap[key]=len(refs)+1;refs.append(dict(key=key,citation=citation,url=url))
def bibclean(s):
 return s.replace('\\&','&').replace('{','').replace('}','').replace('\\l','ł')
addref('lago2021forecasting','Lago, Jesus; Marcjasz, Grzegorz; De Schutter, Bart; Weron, Rafał. Forecasting day-ahead electricity prices: A review of state-of-the-art algorithms, best practices and an open-access benchmark. Applied Energy, 2021, 293:116983.','https://doi.org/10.1016/j.apenergy.2021.116983')
addref('weron2014electricity','Weron, Rafał. Electricity price forecasting: A review of the state-of-the-art with a look into the future. International Journal of Forecasting, 2014, 30(4):1030–1081.','https://doi.org/10.1016/j.ijforecast.2014.08.008')
addref('nowotarski2018recent','Nowotarski, Jakub; Weron, Rafał. Recent advances in electricity price forecasting: A review of probabilistic forecasting. Renewable and Sustainable Energy Reviews, 2018, 81:1548–1568.','https://doi.org/10.1016/j.rser.2017.05.234')
addref('zhang2026dualtimesfield','Zhang, Wencheng; Li, Long; Qin, Huayi; Wu, Zongjuan; Li, Jing; Chen, Wanghu. DualTimesField: Rethinking Time Series as Continuous-Time Trends and Events. ICML, PMLR, 2026, 306:156544–156566.','https://proceedings.mlr.press/v306/zhang26co.html')
addref('aemo_data','Australian Energy Market Operator (AEMO). Aggregated price and demand data. 官方数据入口，访问日期：2026-10-26.','https://www.aemo.com.au/energy-systems/electricity/national-electricity-market-nem/data-nem/aggregated-data')
for b in baselines['references']:addref(b['key'],b['citation'],b.get('url',''))
def cite(keys):return ' '.join('['+str(refmap[k])+']' for k in keys if k in refmap)
def hyperlink(paragraph,url):
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 link=OxmlElement('w:hyperlink');link.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
 run=OxmlElement('w:r');prop=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'1E687E');prop.append(col);size=OxmlElement('w:sz');size.set(qn('w:val'),'17');prop.append(size);run.append(prop);txt=OxmlElement('w:t');txt.text=url;run.append(txt);link.append(run);paragraph._p.append(link)
def clean(s):
 s=str(s).replace('**','').replace('`','')
 for k,n in refmap.items():s=s.replace('['+k+']','['+str(n)+']')
 return s

doc=Document();sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.0)
sec.header_distance=Cm(.9);sec.footer_distance=Cm(.8)
styles=doc.styles
for sn in ['Normal','Body Text']:
 st=styles[sn];st.font.name='Calibri';st.font.size=Pt(10.5);st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.22
for sn,size,col in [('Title',30,'143E52'),('Subtitle',14,'527080'),('Heading 1',18,'143E52'),('Heading 2',13,'1E687E'),('Heading 3',11,'1E687E')]:
 st=styles[sn];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(col);st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(7)
styles['Caption'].font.name='Calibri';styles['Caption'].font.size=Pt(9);styles['Caption']._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');styles['Caption'].font.color.rgb=RGBColor.from_string('527080')
header=sec.header.paragraphs[0];header.text='研究阶段汇报  |  电价尖峰与概率预测';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('2026年10月26日  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)
def p(t,style=None):return doc.add_paragraph(clean(t),style)
def h(t,level=1):return doc.add_heading(clean(t),level)
def chapter(t):
 z=h(t);z.paragraph_format.page_break_before=True;return z
def eq(t):
 z=p(t);z.alignment=WD_ALIGN_PARAGRAPH.CENTER;z.paragraph_format.space_before=Pt(5);z.paragraph_format.space_after=Pt(8)
 for r in z.runs:r.font.name='Cambria';r.font.size=Pt(10)
 return z
def table(headers,rows,widths=None,font=8.7):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=str(x)
 for row in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,row):c.text=clean(x)
 for i,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trPr.append(nosplit)
  if i==0:
   rh=OxmlElement('w:tblHeader');trPr.append(rh)
  for j,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[j]*min(1,16.5/sum(widths)))
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'143E52' if i==0 else ('EEF4F7' if i%2 else 'FFFFFF'));tcPr.append(sh)
   for pa in c.paragraphs:
    if len(rows)<12 and i<len(rows):pa.paragraph_format.keep_with_next=True
    pa.paragraph_format.space_after=Pt(3);pa.paragraph_format.space_before=Pt(3);pa.paragraph_format.line_spacing=1.08
    for r in pa.runs:
     r.font.name="Calibri";r.font.size=Pt(font);r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '172F3B');r.bold=(i==0);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
def fig(file,caption,width=16.6):
 doc.add_picture(str(O/file),width=Cm(width));doc.paragraphs[-1].paragraph_format.keep_with_next=True;p(caption,'Caption')
def label(x):
 return {'baseline_lightgbm':'LightGBM','baseline_lightgbm_cp':'LightGBM+CP','baseline_xgboost':'XGBoost','baseline_xgboost_cp':'XGBoost+CP','baseline_lightgbm_qra':'LightGBM+QRA','baseline_lasso_arx72_cp':'Lasso-ARX72+CP','baseline_patchtst':'PatchTST','baseline_gatets_adapter':'GateTS','baseline_timexer':'TimeXer','baseline_dlinear':'DLinear','baseline_linear':'Linear','baseline_mlp':'MLP','baseline_pricefm':'PriceFM（本地移植）','baseline_ddnn_jsu':'DDNN-JSU','baseline_nhits_qra':'N-HiTS+QRA','tabpfn_v3':'TabPFN-3','chronos2_zeroshot':'Chronos-2','moirai2_zeroshot':'Moirai-2','timesfm25_xreg':'TimesFM-2.5+XReg','Full_joint':'Full_joint（当前方法）','Full_joint_raw':'Full_joint_raw','event_off_same_head':'关闭事件项','time_shift_off_same_head':'关闭时间移动','Anchor_joint':'Anchor_joint','B_retrained':'B_retrained（固定树）'}.get(x,x)
def fmt(m,s=None,n=3):
 if pd.isna(m):return '—'
 return f'{m:.{n}f}'+(f' ± {s:.{n}f}' if s is not None and not pd.isna(s) else '')

p('RESEARCH PROGRESS REPORT','Subtitle')
h('面向电价尖峰的\n概率预测研究\n阶段汇报',0)
p('基线复现 · 趋势—事件建模 · 历史校正与稳健性','Subtitle')
p('\n')
table(['报告信息','内容'],[['汇报日期','2026年10月26日'],['研究对象','澳大利亚 NEM：NSW1、QLD1、TAS1'],['预测任务','过去72小时 → 未来24小时；每小时滚动起报'],['方法版本','Full_joint / asymmetric_balance_20261023'],['证据截至','正式主表更新至2026年10月24日；本报告补充诊断另行标注'],['覆盖范围','19个基线、6个联合方法输出；原始价格与五分位预测']],[3.0,13.5],10)
p('本报告依据工作区的代码、配置、冻结预测与实验记录编制。研究结论以最新同协议结果为主；历史试验用于解释研究路径，未来方案标明为计划。')
chapter('阶段结论与汇报提要')
p('本阶段已完成电价预测公共数据流程、19个基线的适配与统一评估，并将趋势—事件分解思路从直接预测推进到“强树模型基座＋成熟历史检索＋联合校正”。当前方法能够在部分地区和部分误差维度上产生收益，但尚未形成稳定优于最强树基线的总体结果。研究重点应从继续扩展模型数量，转向解释何时值得校正、尖峰误差改善能否覆盖常态时段代价，以及概率区间的校准可靠性。')
p('当前正式结果：Full_joint 宏 MAE 为37.783±0.439，五分位 QS 为24.392±0.188；LightGBM 为37.101±0.015和24.191±0.005。Full_joint 的地区平均 RMSE 从LightGBM的277.064降至275.397，但80%/90%区间覆盖率降至68.56%/82.70%。这些结果表明点误差、极端误差与区间可靠性之间仍存在明显权衡，不能据总体RMSE的下降宣布尖峰问题已解决。')
table(['已完成','尚未解决'],[['19个基线的共同样本评分与完整来源记录','各模型等预算超参数搜索尚未执行'],['当前方法三神经种子、三地区评估','仅神经部分随机变化，尚非整套树＋神经独立重复'],['Anchor平衡损失、非对称尺度、校正限幅等开发实验','尚无方案通过全部预先设定的稳定性门槛'],['机理状态证据门的实现与短流程检查','假设证据不一致，完整确认尚未启动']],[8.2,8.3],9.3)
p('补充尖峰诊断：对同一批冻结预测复算后，Full_joint在真实价格>300及>1000时的MAE较LightGBM分别改善2.10%和0.81%，但常态0–300时段退化4.93%，两个尖峰切片的QS也未改善。下一阶段应重点解决“尖峰点预测收益—常态代价—上尾可靠性”的联合约束。')
h('阅读顺序',2)
for x in ['1 研究背景与任务定义：研究什么、为什么尖峰难预测。','2 实验协议与当前研究阶段：数据、信息边界、评价与完成情况。','3 全部基线方法：原论文机制、引用、本地实现与局限。','4 本研究方法与已尝试路线：从双场到当前联合校正。','5 总体结果与尖峰诊断：19个基线和当前方法完整比较。','6 下一步研究思路：可证伪的假设、对照实验与验收标准。','参考文献与附录：文献更正、结果来源和历史结构对照。']:p(x)
chapter('1 研究背景与任务定义')
h('1.1 研究任务及实际意义',2)
p('本研究关注短期电力现货价格预测：在每个预测起点，利用已观测的历史电价与需求、日历、燃气背景以及当时可获得的预调度和供需预测，输出未来24小时的价格中位数和不确定性区间。预测有助于理解价格风险，并可为后续储能调度、报价与风险管理提供输入；本阶段的直接实验目标是预测质量，而不是已经验证的交易或储能收益。')
p('电价具有周期性、状态变化和厚尾现象。普通时段的价格水平受负荷、燃料成本与可再生出力等影响；少数供需紧张或系统约束事件则可能对应短时高价。低价与负价又具有不同的供需背景。模型既要识别持续趋势，也要避免把偶发事件平滑掉；同时不能为了追逐少量尖峰而在大量普通时段产生错误抬价。相关综述强调应在多个市场、足够长的测试区间和强简单基线上检验新方法。'+cite(['weron2014electricity','lago2021forecasting','nowotarski2018recent']))
h('1.2 本项目的精确定义',2)
eq('输入 Iₜ = {历史72 h价格/需求，日历，起点已知外生预测，合法成熟历史}；输出 Q̂ₜ,ₕ(τ)，h = 0,…,23。')
p('实验地区为NSW1（新南威尔士）、QLD1（昆士兰）和TAS1（塔斯马尼亚），单位为AUD/MWh。当前任务每小时起报一次，预测从起点开始的连续24个交付小时，时间统一按AEST（UTC+10）。它属于滚动多步预测，不能直接等同于某一固定拍卖时刻的一次性日前电价预测。历史实际电价与需求来自AEMO公开数据及本地小时序列。'+cite(['aemo_data']))
p('“尖端电价”在本报告中按研究代码与任务语境统一表述为“电价尖峰”。既有开发代码常以真实价格超过300 AUD/MWh识别正尖峰；本报告同时用超过1000 AUD/MWh作更极端区间的事后诊断，并单独观察负价。阈值是研究分析口径，不代表市场官方统一定义，也不用于重新选择模型。')
h('1.3 文献脉络与本研究切入点',2)
p('本报告中的“全部基线”指当前正式比较表中完成的19个基线配置，而不是文献库中的所有论文。它们覆盖线性与前馈模型、梯度提升树、时序Transformer、分布参数神经网络、校准组合方法和预训练基础模型。组合基线同时引用骨干与校准论文；本地Linear和MLP属于通用对照，不将背景文献冒充其完整模型的原始论文。各篇论文和适配范围在第3章逐项说明。')
p('研究的结构起点是DualTimesField：把时间序列表示成连续趋势场与局部事件场的叠加，以低频参数化表达趋势，以稀疏局部原子表达瞬态。该框架的重建和插值能力为电价建模提供启发，但已知历史曲线的分解成功不等于未知未来尖峰可以被预测。因此，本项目逐步将重点转移到外生状态、成熟历史中的预测误差，以及历史事件能否在当前状态下被可靠采用。'+cite(['zhang2026dualtimesfield']))
h('1.4 当前要回答的研究问题',2)
table(['问题','研究含义'],[['尖峰发生与幅度','在预测起点的信息条件下，是否能提前识别极端高价，而不仅仅重建已知尖峰？'],['尖峰时间位置','历史事件相似时，能否识别其在未来24小时中的位置变化？'],['收益与代价','极端时段的改善是否以普通时段误报、总体MAE或区间失准为代价？'],['跨状态可靠性','历史校正的证据在不同地区、年份与供需状态下是否稳定？']],[3.3,13.2],9.5)
h('2 实验协议与当前研究阶段')
h('2.1 数据、输入与时间划分',2)
table(['项目','当前正式口径'],[['基础数据与时区','NSW1、QLD1、TAS1小时价格/需求；AEST，无夏令时切换'],['静态模型训练','2015-01-01至2021-12-31（各适配器按自身实现使用）'],['选择期','2022年上半年；用于轮数、检查点或既定超参数选择'],['校准期','2022年下半年；CP/QRA等按其固定流程使用'],['正式测试','2023—2024年，每地区17,521个起点，每起点24个标签'],['联合模型神经训练','2018—2021年，使用按时间前推的OOF树预测；选择期2022H1'],['历史记忆','仅使用完整标签已成熟的历史窗口；测试期可因果更新记忆与误差统计'],['随机性','2026/2027/2028；联合方法只改变神经初始化与训练顺序，树固定2026'],['目标价格','原始RRP，identity变换；不以价格截帽后的结果代替正式主表']],[3.3,13.2],9.0)
p('正式测试共17,521×24=420,504个起点—步长单元/地区，每个模型三地区约126万个单元。相邻起点的预测窗口重叠，同一交付小时会被多个提前量预测；这些单元并不是彼此独立的事件样本。主表对每个seed单独评分，先做三地区等权宏平均，再对三个seed计算均值和样本标准差，不先平均预测。')
table(['公共输入块','维度','含义'],[['历史观测','72×2','价格和总需求'],['历史日历','72×10','小时、星期、年内日期、月份等'],['未来日历','24×10','未来交付小时可预先确定的日历'],['起点背景','1','此前7个已结束gas day的燃气价格log背景'],['未来外生预测','24×11','备用容量、预调度价格/需求/可用发电、互联与覆盖标记'],['净负荷预测','24×1','供需平衡的趋势背景'],['容量短缺特征','24×3','相对训练期1%、5%、10%分位阈值的shortfall']],[3.4,2,11.1],8.7)
p('展平后公共输入为1465维。TabPFN使用完整公共输入；Chronos-2、Moirai-2、TimesFM-2.5+XReg受适配器能力约束，仅接收其中部分信息。当前Full_joint还读取因果成熟历史和误差统计，因此它与静态模型具有相同测试样本和公共数据基础，但历史信息范围不同。')
h('2.2 评价指标与汇总方式',2)
eq('MAE = mean|y − q̂₀.₅|；ρτ(u) = u·(τ − 1{u<0})；QS = 2·meanₜ,ₕ,τ ρτ(y − q̂τ)。')
p('公共分位水平为0.05、0.10、0.50、0.90、0.95；评分前按分位轴排序，中位数作为Q组点预测。QS是这五个分位的平均pinball乘2，是离散概率评分代理，不能直接称作对完整分布精确积分的CRPS。RMSE先在各地区按全部预测单元计算，再取三个地区的算术平均。80%/90%覆盖率与区间宽度共同衡量概率输出；覆盖率越大不一定越好，应接近其名义水平并兼顾宽度。')
p('Q组具备完整五分位输出，可比较MAE、QS和两种覆盖率。P组仅比较其原生中位数的点预测指标：Moirai-2和当前TimesFM适配器未提供完整公共五分位，主表中的QS记为“—”，不插值补齐、不与Q组混排概率能力。')
h('2.3 阶段成果与证据边界',2)
p('目前已完成19个基线及当前联合方法6个输出的正式比较，共25个输出×3地区×3seed=225格；旧静态Full/no_fields保留为18格历史结构对照。已有代码与冻结预测的核验、DDNN未截断JSU修复、TimesFM ridge1开发选择和基线数据可用性检查。报告数值以当前CSV和verification.json为准，旧publication.json的21行计数不代表现有主表。')
p('研究仍是回顾性评估：相关年份已被此前实验观察，不应称为独立盲测；当前为固定配置比较，未完成各模型等预算HPO。外生数据已保存的来源时间代理不晚于起点，但真实首次公开发布时间和后续修订尚未全面审计。基础模型的预训练重叠也未审计。上述边界影响外部可复现性和因果解释，不改变本报告对已保存预测数值的客观汇总。')
chapter('3 全部基线论文与方法详解')
p('每项均按“论文思想—本地实现—尖峰预测作用与限制”组织。原论文结果不跨数据集移入本项目总表；第5章全部数字来自本地统一测试。19个基线配置可能共享同一骨干论文，组合方法分别注明校准来源。')
for i,m in enumerate(baselines['models'],1):
 h(f"3.{i} {m['title']}",2)
 p('对应引用：'+cite(m.get('refs',[])))
 for text in m['paragraphs']:p(text)
 if m.get('equations'):
  for x in m['equations']:eq(x)
chapter('4 本研究方法、尝试过程与未解决问题')
fig('architecture.png','图1　当前Full_joint的结构概览。7/28天成熟误差统计进入Anchor；历史分支使用当时样本外树预测构造残差。上图为概览，精确公式见下文。')
for i,s in enumerate(method['sections'],1):
 h(f"4.{i} {s['title']}",2)
 for text in s.get('paragraphs',[]):p(text)
 for x in s.get('equations',[]):eq(x)
 if s.get('table'):
  t=s['table'];table(t['headers'],t['rows'],font=8.6)

chapter('5 全部基线与本研究方法的结果对比')
p('以下是当前25个输出的完整正式结果，价格单位为AUD/MWh。MAE与QS报告三个seed宏平均的均值±样本标准差；RMSE报告地区均值。Q/公共表示完整公共输入、Q/历史表示额外成熟历史、Q/受限和P/受限表示基础模型输入或输出能力受限。数值较小表示误差较低。')
rows=[]
for _,r in main.iterrows():
 group='Q/历史' if 'mature' in r.input_scope else ('P/受限' if r.table=='P' else 'Q/受限' if 'restricted' in r.input_scope else 'Q/公共')
 rows.append([label(r.model),group,fmt(r.mae_mean,r.mae_std),fmt(r.quantile_score_proxy_mean,r.quantile_score_proxy_std),fmt(r.rmse_pooled_mean)])
table(['模型/输出','组别','宏 MAE ↓','宏 QS ↓','RMSE ↓'],rows,[4.6,1.6,3.5,3.5,2.7],8.0)
p('表1　全量主表。原始来源：main_comparison_20261023/main.csv，最新维护至2026-10-24。B_retrained是同一固定树的重复参照；基础模型的0标准差可能来自确定性重复，不能据此推断统计稳定性优于随机训练模型。CP的点预测未改动，因此其MAE与对应树基线相同。','Caption')
fig('comparison_mae.png','图2　19个基线与当前Full_joint的总体MAE。颜色突出LightGBM与本研究方法；输入能力差异见表1，图中排序不构成公平HPO后的算法优劣结论。')
h('5.1 总体结果能够支持什么',2)
p('LightGBM是当前19个外部基线中总体MAE与QS最强的参照。Full_joint较LightGBM的MAE增加0.681（约1.84%），QS增加0.201（约0.83%），地区平均RMSE降低1.667（约0.60%）。相对XGBoost，Full_joint的MAE减少0.147（约0.39%）、QS减少1.230（约4.80%）；因此它已接近强树基线，但尚未超过最强参照。')
p('在当前预算下，复杂时序网络或基础模型没有自动取得更好的总体结果。该发现仅适用于本数据、固定配置和适配范围：例如PriceFM采用本地移植而非其原论文完整预训练流程，三种时序基础模型的输入受限，不能将主表排名概括成“基础模型无效”。')
p('LightGBM+QRA的地区平均RMSE为269.466，优于Full_joint的275.397，但其MAE和QS更高。各模型对中位数误差、极端偏差和分布质量的优化方向不同，需要多指标判断；不能只挑选某一个有利指标作为综合胜出依据。')
h('5.2 地区差异：QLD1是主要瓶颈',2)
rr=[]
for region in ['NSW1','QLD1','TAS1']:
 for m in ['baseline_lightgbm','baseline_xgboost','Full_joint']:
  r=reg[(reg.model==m)&(reg.region==region)].iloc[0];rr.append([region,label(m),fmt(r.mae_mean,r.mae_std),fmt(r.quantile_score_proxy_mean,r.quantile_score_proxy_std),f'{r.coverage_90_mean*100:.2f}%'])
p('表2　三地区的核心模型比较','Caption').paragraph_format.keep_with_next=True
table(['地区','模型','MAE ↓','QS ↓','90%覆盖率'],rr,[1.5,4.3,3.6,3.6,2.5],8.3)
fig('regional_comparison.png','图3　三地区的MAE与QS。Full_joint在NSW1改善较一致；QLD1两项主要评分均退化，不能以跨地区平均掩盖此差异。')
p('相对LightGBM，Full_joint在NSW1的MAE由44.333降至44.036，QS由31.318降至30.854；在QLD1，MAE由38.213升至40.746，QS由24.919升至25.956；在TAS1，MAE小幅下降而QS略升。QLD1的MAE种子标准差达到1.461，是下一步定位校正不稳定性的重点。')
h('5.3 区间校准与消融解释',2)
fig('coverage.png','图4　概率区间覆盖率。虚线为80%和90%名义水平。Full_joint区间偏窄、覆盖不足；CP则显著扩大区间，覆盖偏高但QS变差。')
p('Full_joint的80%覆盖率为68.56%，低于LightGBM的72.50%；90%覆盖率为82.70%，低于85.10%。平均90%区间宽度由141.35缩至118.08。因此“区间更窄”没有同时获得更好的可靠性。QLD1尤其明显：80%覆盖率从78.08%降至66.73%，提示Anchor尺度头及后续校正需先解决条件校准。')
p('Full_joint相对同检查点Anchor_joint将MAE从37.908降至37.783、QS从24.467降至24.392，说明历史分支在该输出上有局部增益；但关闭事件项的MAE为37.754，反而略低于完整方法。关闭时间移动后的QS为24.384，也略低于24.392。它们是同检查点干预而非独立重训消融，目前不足以证明事件分支或时间移动具有稳定必要性。')
h('5.4 尖峰与常态分层：本报告补充诊断',2)
# Filled from independently audited files below.
if (O/'tail_report_payload.json').exists():
 tail=json.loads((O/'tail_report_payload.json').read_text())
 for x in tail.get('paragraphs',[]):p(x)
 for t in tail.get('tables',[]):
  if t.get('title'):p(t['title'],'Caption').paragraph_format.keep_with_next=True
  table(t['headers'],t['rows'],font=8.1)
else:
 p('当前正式主表提供全时段汇总；本报告对保存预测的额外尖峰分层核验结果见附带结果审计文件。分层以实际价格作事后条件，不能在预测时读取该条件。')
fig('tail_tradeoff.png','图5　Full_joint相对LightGBM的分层MAE变化。负值为改善，正值为退化；尖峰收益与常态代价同时存在。')
p('尖峰点误差改善尚未转化为尖峰概率预测改善：高价QS从808.368升至815.337，极端QS从2983.550升至3034.190。这与历史分支仅共同平移五分位、不能单独扩张上尾的能力边界相呼应，但尚不足以单独确定失效原因。')
h('5.5 结论边界',2)
p('本阶段的可信结论是：强树基线仍难以全面超越，联合历史校正已有局部收益，但事件机制的增量价值和跨地区稳定性尚未得到充分证实。选择期guarded规则及alpha回退只能约束选择期表现，不保证未来改善；本次Full_joint没有回退格也不是成功的充分证据。没有依赖配对时间块的显著性分析和未触碰新测试区间之前，不将小幅差异表述为统计显著提升。')

chapter('6 下一步研究思路与实验安排')
p('下一阶段建议围绕“证据可靠时校正、不可靠时保持强基线”的可检验目标，先分析现有误差，再增加结构。优先确保极端时段收益、普通时段代价和区间校准同时可追溯；各项方案均应在独立开发目录冻结配置，不依据2023—2024测试排名反复选优。')
# next_steps may be structured records or strings
for i,s in enumerate(method.get('next_steps',[]),1):
 if isinstance(s,str):p(s)
 elif isinstance(s,dict):
  h(f"6.{i} "+s.get('title',s.get('name',s.get('question',s.get('hypothesis','研究实验')))),2)
  for k,v in s.items():
   if k in ['title','name','question']:continue
   names={'hypothesis':'待检验假设','experiment':'实验设置','controls':'对照设置','metrics':'评价指标','acceptance':'验收标准','priority':'优先级','risk':'风险','stop_rule':'停止规则','status':'当前状态','design':'实验设计','matrix':'实验矩阵','decision':'判断规则'}
   if isinstance(v,list):v='；'.join(str(x) for x in v)
   elif isinstance(v,dict):v='；'.join(f'{a}：{b}' for a,b in v.items())
   p(names.get(k,k)+'：'+str(v))
h('建议采用的阶段里程碑',2)
table(['顺序','交付物','进入下一步的条件'],[['第一步','当前预测的分层诊断：地区×季节×步长×价格带；以日或事件为块的置信区间','确定QLD1退化来自位置、尺度、普通时段误报或事件错位，避免凭单一均值猜测'],['第二步','强树基座、Anchor-only、趋势-only、事件-only、Full的独立重训对照','相同历史范围与训练预算，证明新增分支增益，保留失败结果'],['第三步','机理证据门 vs 同平均收缩固定门 vs 简单限幅','验证收益来自状态证据而非单纯减小校正量；若假设不支持则停止扩展'],['第四步','冻结后的跨年份/跨地区确认与可靠性评估','宏MAE/QS不退化，尖峰指标改善，普通时段误差及条件覆盖不出现不可接受代价']],[1.5,6.6,8.4],9)
p('对于尚未观察的新时间段，先冻结时间边界、尖峰阈值、候选集合、选择规则和验收标准。若现有研究已多次使用2023—2024年结果，应将其保留为回顾性参照，另留未触碰时段作确认；若暂时没有新数据，则明确只报告开发结论，不包装成独立外推能力。')

chapter('参考文献')
p('以下文献编号与第1—4章对应。原论文、方法组件和官方实现引用分列；组合基线并不等同于完整复现相关比较论文。网页核对日期为2026年10月26日。')
for i,r in enumerate(refs,1):
 pa=p(f"[{i}] {r['citation']}");pa.paragraph_format.keep_with_next=True
 pa.paragraph_format.space_after=Pt(4)
 if r.get('url'):
  pa=p('');hyperlink(pa,r['url']);pa.paragraph_format.space_after=Pt(6)
  for run in pa.runs:run.font.size=Pt(8.5);run.font.color.rgb=RGBColor.from_string('1E687E')

chapter('附录A 文献与实验来源说明')
for c in baselines.get('corrections',[]):p(c if isinstance(c,str) else c.get('correction',str(c)))
p('正文结果直接取自当前主表，早期paper/manuscript中季度重训、价格下界、截帽或seed平均预测的成绩不混入本报告正式表。Informer/iTransformer保留于历史代码，但不在当前已完成19基线中；它们没有被作为现行主表缺失项补造数值。')
paths=[('forecasting/CITATIONS.md','19个基线与引用索引'),('forecasting/README.md','模型身份、实现入口与当前版本'),('docs/v5_latest_aligned_20261023.md','当前Full_joint冻结协议'),('logs/forecasting/main_comparison_20261023/main.csv','25个输出的正式宏平均主表'),('logs/forecasting/main_comparison_20261023/regional.csv','三地区结果'),('logs/forecasting/main_comparison_20261023/verification.json','来源预测、评分和完整性核验'),('logs/forecasting/v5_calibration_robustness_20261023/README.md','校正稳健性开发实验'),('logs/forecasting/v5_mechanism_state_20261023/README.md','机理状态诊断与短流程'),('logs/forecasting/baseline_maintenance_20261024/README.md','数据可用性、TimesFM与TabPFN维护')]
for f,desc in paths:p(desc+'：'+f)
p('上述路径均相对于工作区 /root/workspace/Code/DualTimesField。报告配套目录保存编制脚本、图表、机器可读方法说明与结果审计，可追踪每个结论。')
h('附录B 历史结构对照（同测试评分，非当前方法）')
hist=pd.read_csv(D/'historical_controls.csv')
rows=[]
for _,r in hist.iterrows():rows.append([r.model,fmt(r.mae_mean,r.mae_std),fmt(r.quantile_score_proxy_mean,r.quantile_score_proxy_std),fmt(r.rmse_pooled_mean)])
table(['历史模型','MAE','QS','RMSE'],rows,[5,4,4,3.5],8.8)
p('这些行记录旧静态Full/no_fields的结构效果，不能替代新版Full_joint。历史版本、开发年份和最终测试具有各自用途；本报告统一保留其身份，以便阶段评审判断研究迭代是否获得实质进展。')

chapter('附录C 全部输出的尖峰分层MAE')
p('本报告新增事后诊断；各地区先评分、等权宏平均后对seed平均。高价与极端分组重叠。完整标准差、QS、偏差和覆盖率保存在latest_spike_diagnostics_macro.csv。')
taildf=pd.read_csv(O/'latest_spike_diagnostics_macro.csv').set_index(['model','band'])
rows=[]
for model_id in main.model:
 rows.append([label(model_id)]+[f'{taildf.loc[(model_id,b),"mae_mean"]:.3f}' for b in ['negative','normal_0_to_300','high_gt300','extreme_gt1000']])
table(['模型/输出','负价 <0','常态0–300','高价 >300','极端 >1000'],rows,[5.1,2.85,2.85,2.85,2.85],8.3)
out=O/'电价尖峰预测研究_阶段汇报_20261026.docx';doc.save(out)
manifest={'generated_date':'2026-10-26','docx':str(out),'main_rows':len(main),'reference_count':len(refs),'baseline_count':len(baselines['models']),'sources':{str(D/f):hashlib.sha256((D/f).read_bytes()).hexdigest() for f in ['main.csv','regional.csv','per_seed.csv','verification.json']},'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)}
(O/'report_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps(manifest,ensure_ascii=False,indent=2))
