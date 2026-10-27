from pathlib import Path
P=Path(__file__).parent
# Reuse the already audited data loading, bibliography and model-name mappings.
source=(P/'build_report.py').read_text()
exec(compile(source[:source.index("p('RESEARCH PROGRESS REPORT'")],str(P/'build_report.py'),'exec'))
from docx.enum.text import WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from copy import deepcopy
B=P/'baselines_academic.json'
if B.exists():baselines=json.loads(B.read_text())

ordered_keys=['weron2014electricity','nowotarski2018recent','lago2021forecasting','zhang2026dualtimesfield']
model_order=['linear','dlinear','mlp','lasso_arx72_cp','lightgbm','xgboost','patchtst','gatets_adapter','timexer','pricefm','ddnn_jsu','lightgbm_cp','xgboost_cp','lightgbm_qra','nhits_qra','tabpfn_v3','chronos2_zeroshot','moirai2_zeroshot','timesfm25_xreg']
for model_id in model_order:
 for key in next(m for m in baselines['models'] if m['id']==model_id)['refs']:
  if key not in ordered_keys:ordered_keys.append(key)
ordered_keys.append('aemo_data')
by_key={r['key']:r for r in refs};refs=[by_key[k] for k in ordered_keys];refmap={r['key']:i+1 for i,r in enumerate(refs)}

doc=Document();styles=doc.styles
for st in styles:
 if st.type==1:
  st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0)
  st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
normal=styles['Normal'];normal.paragraph_format.line_spacing=Pt(22);normal.paragraph_format.first_line_indent=Pt(24);normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.widow_control=True
for n,sz in [('Title',22),('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
 st=styles[n];st.font.name='Times New Roman';st.font.size=Pt(sz);st.font.bold=True;st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(8);st.paragraph_format.keep_with_next=True
styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Caption'].font.size=Pt(10.5);styles['Caption'].font.italic=False;styles['Caption'].font.bold=False;styles['Caption'].paragraph_format.first_line_indent=Pt(0);styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;styles['Caption'].paragraph_format.line_spacing=1.15;styles['Caption'].paragraph_format.space_before=Pt(5);styles['Caption'].paragraph_format.space_after=Pt(6)
for s in ['TOC Heading','TOC 1','TOC 2']:
 if s in styles:
  styles[s].paragraph_format.first_line_indent=Pt(0)

def setup(sec,kind):
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.5);sec.bottom_margin=Cm(2.4);sec.left_margin=Cm(2.7);sec.right_margin=Cm(2.4);sec.header_distance=Cm(1.25);sec.footer_distance=Cm(1.25)
 sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
 sec.header.paragraphs[0].text='';sec.footer.paragraphs[0].text=''
 if kind!='cover':
  hp=sec.header.paragraphs[0];hp.text='面向电价尖峰的概率预测方法研究';hp.alignment=WD_ALIGN_PARAGRAPH.CENTER;hp.paragraph_format.first_line_indent=Pt(0)
  for r in hp.runs:r.font.size=Pt(9)
  f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.paragraph_format.first_line_indent=Pt(0);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
  pg=OxmlElement('w:pgNumType');pg.set(qn('w:fmt'),'lowerRoman' if kind=='front' else 'decimal');pg.set(qn('w:start'),'1');sec._sectPr.append(pg)
setup(doc.sections[0],'cover')
chapter_num=0;fig_num=0;tab_num=0;eq_num=0

def polish(text):
 text=clean(text)
 subs={'本报告':'本文','本项目':'本研究','起报':'预测起点','本地实现':'本文实现','正式主表':'总体比较','主表':'结果表','不能写成':'不同于','不能称为':'不构成','不能宣布':'尚不足以认定','不能证明':'尚不足以证明','不应称为':'不属于','本轮':'该实验','当前完整配置下':'完整配置下','代码':'实现','P08':'相关比较研究','P16':'相关校准研究','J34':'相关比较研究'}
 for a,b in subs.items():text=text.replace(a,b)
 return text

def p(text,style=None,indent=True):
 pa=doc.add_paragraph(style=style);pa.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 if not indent:pa.paragraph_format.first_line_indent=Pt(0)
 for part in re.split(r'(\[\d+(?:[,–\-]\d+)*\])',polish(text)):
  if not part:continue
  r=pa.add_run(part)
  if re.fullmatch(r'\[\d+(?:[,–\-]\d+)*\]',part):r.font.superscript=True;r.font.size=Pt(9)
 return pa

def h(title,level=2):return doc.add_heading(title,level)
def chapter(n,title,first=False):
 global chapter_num,fig_num,tab_num,eq_num
 chapter_num=n;fig_num=tab_num=eq_num=0
 hd=h(f'第{n}章  {title}',1)
 if not first:hd.paragraph_format.page_break_before=True
 return hd

equation_elements=Document(P/'academic_equations.docx')._element.xpath('//m:oMath')
equation_index=0
def eq(text):
 global eq_num,equation_index
 eq_num+=1
 t=doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Cm(14.1);t.columns[1].width=Cm(1.7)
 for i,w in enumerate([14.1,1.7]):t.cell(0,i).width=Cm(w)
 pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for key in ['top','bottom','left','right','insideH','insideV']:
  el=OxmlElement('w:'+key);el.set(qn('w:val'),'nil');borders.append(el)
 pr.append(borders);t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 for c in t.rows[0].cells:
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for pa in c.paragraphs:pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.line_spacing=1.05;pa.paragraph_format.space_before=Pt(5);pa.paragraph_format.space_after=Pt(5)
 pa=t.cell(0,0).paragraphs[0];pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
 equation_index+=1;pa.add_run().add_picture(str(P/f'equation_{equation_index:02d}.png'),width=Cm(13.8))
 num=t.cell(0,1).paragraphs[0];num.alignment=WD_ALIGN_PARAGRAPH.RIGHT;num.add_run(f'（{chapter_num}-{eq_num}）').font.size=Pt(10.5)
 return pa

def caption(text):
 pa=doc.add_paragraph(text,'Caption');return pa

def table(title,headers,rows,widths=None,font=10):
 global tab_num
 tab_num+=1;cap=caption(f'表{chapter_num}-{tab_num}  {title}');cap.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if not widths:widths=[15.8/len(headers)]*len(headers)
 widths=[w*15.8/sum(widths) for w in widths]
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=str(x)
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=polish(str(x))
 pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for k in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+k);e.set(qn('w:val'),'single' if k in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'10');e.set(qn('w:color'),'000000');borders.append(e)
 pr.append(borders)
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if i==0:
    cb=OxmlElement('w:tcBorders');bt=OxmlElement('w:bottom');bt.set(qn('w:val'),'single');bt.set(qn('w:sz'),'6');cb.append(bt);c._tc.get_or_add_tcPr().append(cb)
   for pa in c.paragraphs:
    pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.line_spacing=1.1;pa.paragraph_format.space_before=Pt(2);pa.paragraph_format.space_after=Pt(2)
    if len(rows)<=10 and i<len(rows):pa.paragraph_format.keep_with_next=True
    for r in pa.runs:r.font.name='Times New Roman';r.font.size=Pt(font);r.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

def fig(file,title,width=15.5):
 global fig_num
 fig_num+=1;doc.add_picture(str(P/file),width=Cm(width));pa=doc.paragraphs[-1];pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.line_spacing=1.0;pa.paragraph_format.keep_with_next=True;caption(f'图{chapter_num}-{fig_num}  {title}')

def chapter_end(num,text):h(f'{num} 小结',2);p(text)

# Cover
pa=p('研究阶段报告',indent=False);pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.space_before=Pt(100)
for r in pa.runs:r.font.size=Pt(18)
pa=doc.add_paragraph('面向电价尖峰的概率预测方法研究','Title');pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.space_before=Pt(40);pa.paragraph_format.space_after=Pt(22)
pa=p('基于趋势—事件分解与成熟历史校正',indent=False);pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in pa.runs:r.font.size=Pt(15)
pa=p('2026年10月',indent=False);pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.space_before=Pt(145)
setup(doc.add_section(WD_SECTION_START.NEW_PAGE),'front')
h('摘  要',1)
p('电力现货价格同时具有周期性、非平稳性与厚尾特征，少量短时尖峰对整体风险产生显著影响。现有模型在降低总体平均误差与刻画极端价格之间仍存在权衡，单纯提高历史曲线重建精度亦难以保证未来尖峰预测能力。针对上述问题，本研究以澳大利亚国家电力市场三个地区的小时数据为对象，研究利用供需预测信息和成熟历史残差改善滚动电价概率预测的方法。')
p('研究构建了过去72小时输入、未来24小时输出的统一实验流程，完成19个基线模型的适配与评估。在方法设计上，以冻结的LightGBM分位数模型为基础，利用近期成熟误差进行非对称位置—尺度校正，并将历史曲线分解为低频趋势与局部事件。随后，通过条件检索、事件时间对齐与采用门生成有界历史残差，采用联合损失训练神经模块，并在选择期固定检查点及校正收缩系数。')
p('在2023—2024年测试区间内，所研究方法的宏平均MAE和五分位QS分别为37.783和24.392，较LightGBM分别上升1.84%和0.83%。分层结果表明，该方法在实际价格超过300和1000 AUD/MWh时的MAE分别降低2.10%和0.81%，而在0—300 AUD/MWh常态时段上升4.93%，且高价区间的概率评分未同步改善。上述结果说明，成熟历史校正具有局部尖峰点预测价值，但其总体收益受到常态误校正与区间可靠性的限制。后续研究将围绕校正证据、事件发生与强度解耦及跨时期稳定性展开，以建立兼顾常态精度和极端风险的电价预测方法。')
pa=p('关键词：电价预测；概率预测；电价尖峰；趋势—事件分解；历史残差校正',indent=False);pa.paragraph_format.space_before=Pt(12)
hd=h('目 录',1);hd.paragraph_format.page_break_before=True
ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),'9');hd._p.get_or_add_pPr().append(ol)
pa=doc.add_paragraph();pa.paragraph_format.first_line_indent=Pt(0)
for tag,text in [('begin',None),('instr',r'TOC \o "1-2" \h \z \u'),('separate',None),('text','目录将在文档更新时生成。'),('end',None)]:
 r=OxmlElement('w:r')
 if tag in ['begin','separate','end']:
  x=OxmlElement('w:fldChar');x.set(qn('w:fldCharType'),tag)
 elif tag=='instr':x=OxmlElement('w:instrText');x.set(qn('xml:space'),'preserve');x.text=text
 else:x=OxmlElement('w:t');x.text=text
 r.append(x);pa._p.append(r)
setup(doc.add_section(WD_SECTION_START.NEW_PAGE),'body')
chapter(1,'绪论',True)
h('1.1 研究背景与意义')
p('电价预测是电力市场运行与参与者风险管理的重要基础。随着可再生能源出力、负荷水平和系统约束随时间变化，现货价格不仅呈现周期规律，也可能出现短时高价与负价。对预测模型而言，常态价格决定多数样本上的误差水平，稀少尖峰则显著影响极端损失与风险判断。因此，需要同时关注点预测精度、概率分布质量与不同价格状态下的可靠性。'+cite(['weron2014electricity','nowotarski2018recent']))
p('已有综述指出，电价预测方法应在多个市场与较长测试区间内，结合强统计和机器学习基线进行验证，以避免有限样本和评价口径差异导致的结论偏差。'+cite(['lago2021forecasting'])+'本研究据此构建共同样本、共同价格尺度和共同评分下的比较流程，将复杂模型的增量价值落实到可复核的数值结果。')
h('1.2 研究问题与技术挑战')
p('本研究关注每小时滚动起点下的24小时概率预测。在实际价格尚不可见时，模型需要依据此前72小时的价格与需求以及起点已知外生信息，识别潜在尖峰的发生、位置和强度。由于尖峰样本稀少且与供需状态相关，直接优化总体损失的模型可能倾向于平滑极端值；相反，过度强调尖峰又可能在常态时段产生错误抬价。')
p('另一项挑战是历史模式的可迁移性。相似的历史曲线并不意味着当前市场具有相同价格形成条件，历史重建精确也不等价于未来预测有效。若校正缺乏充分的状态支持，趋势与事件分支可能同时产生偏移，最终抵消强基线已经获得的常态精度。此外，中位数校正和上尾不确定性具有不同目标，仅对所有分位数共同平移并不能充分表达尖峰风险。')
h('1.3 研究内容与阶段进展')
p('围绕上述问题，本阶段的研究包括基线体系构建、趋势—事件方法探索和预测误差分析三个方面。首先，完成19个基线配置的统一适配，覆盖线性模型、梯度提升树、深度时序网络、概率校准组合以及预训练基础模型。其次，以DualTimesField的趋势与事件分离思想为出发点，逐步探索直接分解、历史检索和强基座残差校正，并形成具有非对称Anchor头的联合方法。'+cite(['zhang2026dualtimesfield']))
p('最后，在三个地区和三个种子标签下比较当前方法及其干预输出，并进一步按实际价格区间分解误差。现阶段已观察到部分尖峰点误差改善，但尚未形成稳定超越LightGBM的总体结果。研究进展体现为明确了历史校正的作用范围及主要失效模式，为后续可靠性约束提供了实验依据。')
h('1.4 章节安排')
p('第2章介绍相关研究与全部基线方法，说明原始论文机制及本文适配方式。第3章给出问题定义、数据设置与评价协议。第4章阐述方法演进及当前联合模型。第5章比较总体、地区、概率区间和尖峰分层结果，并分析干预与稳健性实验。第6章总结阶段结论，讨论现有不足与后续研究方向。')

chapter(2,'相关研究与基线方法')
p('电价预测研究既包含针对市场数据设计的统计与机器学习方法，也吸收了通用时间序列预测和预训练模型的进展。本文选取19个已完成实验的基线配置，分别从线性结构、非线性阈值、时序表征、分布建模和迁移学习等角度检验联合方法的增量价值。组合模型的骨干和校准机制分别引用其来源文献，所有数值比较均采用本文实验结果。')
groups=[('统计与前馈预测模型',['linear','dlinear','mlp','lasso_arx72_cp']),('梯度提升树模型',['lightgbm','xgboost']),('深度时序表征模型',['patchtst','gatets_adapter','timexer','pricefm']),('概率分布与预测校准方法',['ddnn_jsu','lightgbm_cp','xgboost_cp','lightgbm_qra','nhits_qra']),('预训练基础模型',['tabpfn_v3','chronos2_zeroshot','moirai2_zeroshot','timesfm25_xreg'])]
model_map={m['id']:m for m in baselines['models']}
for g,(name,ids) in enumerate(groups,1):
 h(f'2.{g} {name}')
 for j,modelid in enumerate(ids,1):
  m=model_map[modelid];title=m['title'].replace('：','与',1) if False else m['title']
  h(f'2.{g}.{j} {title}',3)
  for k,text in enumerate(m['paragraphs']):p(text+(cite(m['refs']) if k==0 else ''))
h('2.6 现有方法的局限与研究切入点')
p('线性与前馈方法提供了结构简洁的参照，梯度提升树能够有效利用展平后的供需与预调度特征，时序网络则强调局部模式和跨时间依赖。概率校准方法改善区间构造的灵活性，但其实际表现仍受校准样本分布与模型误差结构影响。基础模型具有跨数据迁移潜力，然而输入能力、原生分位数范围及预训练数据边界使其在本文任务中的比较需要分别解释。')
p('这些方法的共同局限在于，整体预测损失与稀有尖峰表现并不必然一致。本文进一步检验成熟历史中的趋势残差和事件残差能否在当前供需条件下提供增量信息，并通过受限校正降低对强基座常态预测的干扰。')

chapter(3,'问题定义与实验设置')
h('3.1 滚动多步概率预测任务')
p('设t为预测起点，历史长度L=72，预测长度H=24。输入信息包括历史电价与需求、已知日历、当时可用的供需和预调度预测，以及在相应模型中被允许读取的成熟历史。目标为预测未来24个交付小时的条件分位数：')
eq('Q̂ₜ,ₕ(τ) = fθ(Iₜ, h, τ),  h = 0,…,23,  τ ∈ {0.05,0.10,0.50,0.90,0.95}')
p('预测步长h=0对应从t开始的交付小时。实验每小时更新一次预测起点，时间统一为AEST（UTC+10），因此该任务属于滚动多步预测，与固定拍卖时刻的一次性日前预测具有不同信息边界。价格单位为AUD/MWh，中位数分位q̂₀.₅作为点预测。')
h('3.2 数据来源与输入特征')
p('实验使用澳大利亚国家电力市场NSW1、QLD1和TAS1三个地区的数据，分别对应新南威尔士、昆士兰和塔斯马尼亚。实际价格与需求取自AEMO公开数据及其本地小时序列；外生信息包括燃气背景、预调度价格与需求、备用容量、净负荷以及互联相关预测。'+cite(['aemo_data']))
table('公共输入特征及其维度',['特征块','维度','变量含义'],[['历史观测','72×2','电价与总需求'],['历史日历','72×10','小时、星期、年内日期、月份等'],['未来日历','24×10','交付小时可预先确定的日历变量'],['起点背景','1','此前7个已结束gas day的燃气价格对数背景'],['未来外生预测','24×11','备用容量、预调度及互联预测、数据覆盖标记'],['净负荷预测','24×1','供需平衡的趋势背景'],['容量短缺特征','24×3','相对训练期1%、5%、10%分位阈值的shortfall']],[3,2,10.8])
p('公共输入展平后为1465维。TabPFN使用完整公共输入，其他三个时序基础模型仅接收其接口所支持的部分变量。当前联合方法进一步读取成熟历史与误差统计，其历史信息范围大于静态基线，因而共同测试样本不意味着完全相同的信息权限。')
h('3.3 时间划分与训练协议')
table('训练、选择、校准与测试区间',['阶段','时间范围及用途'],[['静态训练','2015—2021年；各模型按固定实现使用训练数据'],['选择','2022年上半年；确定检查点、轮数或既定选择参数'],['校准','2022年下半年；供CP/QRA等方法按预定流程使用'],['正式测试','2023—2024年；每地区17,521个预测起点'],['联合模型训练','2018—2021年；使用按时间前推生成的OOF树预测'],['历史更新','完整24小时标签成熟后，才可纳入历史记忆与误差统计']],[3,12.8])
p('每地区测试包含17,521×24=420,504个起点—步长评分单元。由于窗口重叠，同一交付小时可能对应多个预测提前量，评分单元之间存在相关性。三个种子标签为2026、2027和2028；联合方法仅改变神经初始化和训练顺序，树种子固定为2026。确定性模型的重复结果以及固定树参照不代表独立随机训练。')
p('各模型采用固定配置，未进行统一计算预算下的超参数搜索。当前联合模型训练20轮，全部拟合与检查点选择完成后再评估测试区间；其raw输出与guarded输出采用预先规定的独立选择规则。')
h('3.4 评价指标')
p('本文采用平均绝对误差MAE衡量点预测，采用五分位评分QS评价概率输出。定义预测误差u=y−q̂τ，分位损失及指标为：')
eq('ρτ(u) = u[τ − 1(u < 0)]')
eq('MAE = mean |y − q̂₀.₅|,    QS = 2 meanₜ,ₕ,τ ρτ(y − q̂τ)')
p('QS为公共五个分位水平上的离散评分代理，并非完整分布上的精确CRPS积分。评分前按分位轴排序，Q组使用排序后的中位数；P组仅使用其原生中位数计算点误差。Moirai-2和TimesFM适配器不具备完整公共五分位输出，因此不填补其QS。')
p('RMSE在各地区全部预测单元上计算后取算术平均。80%区间由q̂₀.₁和q̂₀.₉构成，90%区间由q̂₀.₀₅和q̂₀.₉₅构成；覆盖率应与区间宽度联合解释。每个seed先独立评分，再对三个地区等权平均，最后计算三个seed的均值与样本标准差。')
h('3.5 尖峰划分与比较边界')
p('价格状态分为负价y<0、常态0≤y≤300以及高价y>300，并在高价内部单独考察极端价格y>1000。高价与极端组存在包含关系，这些阈值属于研究诊断口径，并非市场统一定义。分组只使用实际价格作事后分析，不用于预测时路由或重新选择模型。')
p('该实验属于回顾性比较。外生归档中已保存的来源时间代理均未晚于预测起点，但首次公开时间与后续修订尚未全面审计；预训练基础模型的数据重叠亦未完成审计。因此，当前结果主要支持所保存数据和固定协议下的性能比较，不构成严格实时无泄露或独立盲测的认证。')

chapter(4,'基于成熟历史校正的趋势—事件联合预测方法')
h('4.1 设计动机与方法演进')
p('本研究的初始假设是，电价可分解为缓慢变化的趋势与局部事件，两者应通过不同表示建模。原始DualTimesField采用低频趋势场与局部事件原子，其对已知曲线的重建能力为该假设提供了表示层面的依据。'+cite(['zhang2026dualtimesfield'])+'在预测任务中，关键问题进一步转变为：历史分解结果中哪些信息能够在当前市场状态下被可靠迁移。')
for text in method['sections'][1]['paragraphs']:
 # Method history is kept as narrative, with detailed numerical comparisons reserved for experiments.
 p(text)
h('4.2 联合模型总体结构')
p('当前方法由冻结的分位数树基座、非对称Anchor校正和趋势—事件历史残差三个部分构成。基座负责利用全部公共输入形成初始预测；Anchor调整位置与上下尾尺度；历史分支在成熟记录中检索条件相关的误差结构，并生成有界修正。神经模块在统一损失下联合训练，基础树独立拟合且保持冻结。')
fig('architecture.png','趋势—事件联合预测模型的总体结构')
h('4.3 基础分位预测与非对称Anchor校正')
for text in method['sections'][2]['paragraphs'][:2]:p(text)
for e in method['sections'][2]['equations'][:2]:eq(e)
p('其中Bτ为基础模型分位数，Aτ为Anchor输出，δ控制条件位置偏移，s−和s+分别控制下尾与上尾尺度。非对称设计使上下尾不必共享相同收缩程度，而恒等初始化保留初始树预测。')
h('4.4 成熟历史的趋势—事件表示')
for text in method['sections'][2]['paragraphs'][2:]:p(text)
for e in method['sections'][2]['equations'][2:]:eq(e)
p('其中LP表示低通趋势提取。历史残差采用该记录形成时的样本外基础预测作为参照，使残差反映当时真实预测误差，减少以已见标签回算预测产生的偏差。')
h('4.5 条件检索与事件时间对齐')
for text in method['sections'][3]['paragraphs'][:2]:p(text)
for e in method['sections'][3]['equations'][:2]:eq(e)
p('wT,j与wE,j分别表示趋势和事件的候选权重，gT与gE控制两类修正的采用程度。R对所有分位数作共同平移，因而保持Anchor的区间宽度。这种结构有利于分离位置补偿和尺度调整，但也限制了历史事件分支对极端上尾的独立表达能力。')
h('4.6 联合损失与模型选择')
for text in method['sections'][3]['paragraphs'][2:]:p(text)
for e in method['sections'][3]['equations'][2:]:eq(e)
p('损失中的重建项约束历史表示，门值正则限制历史过度采用，Anchor直接监督与相对基座移动惩罚共同抑制不必要的校正。选择期守护仅提供有限样本上的保守选择，不构成对未来误差的保证。')
h('4.7 本章小结')
p('本章建立了从强树基座到非对称校正，再到成熟历史残差迁移的预测流程。该方法试图把历史趋势与事件作为有条件采用的增量信息，而不是强制叠加的先验。其有效性取决于校正证据的跨状态稳定性，以及尖峰收益能否覆盖常态代价；相关问题将在实验中检验。')

chapter(5,'实验结果与分析')
h('5.1 总体预测性能')
p('表5-1汇总19个基线与联合方法6个输出的完整结果，共25个输出、225个模型—地区—种子组合。Q/公共表示完整公共输入，Q/历史表示额外成熟历史，Q/受限与P/受限表示基础模型输入或输出能力受限。MAE和QS报告均值±样本标准差，RMSE为地区均值；所有价格指标单位为AUD/MWh。')
rows=[]
for _,r in main.iterrows():
 group='Q/历史' if 'mature' in r.input_scope else ('P/受限' if r.table=='P' else 'Q/受限' if 'restricted' in r.input_scope else 'Q/公共')
 rows.append([label(r.model),group,fmt(r.mae_mean,r.mae_std),fmt(r.quantile_score_proxy_mean,r.quantile_score_proxy_std),fmt(r.rmse_pooled_mean)])
table('全部模型的总体预测结果',['模型/输出','组别','MAE↓','QS↓','RMSE↓'],rows,[4.5,1.5,3.5,3.5,2.5],9)
p('LightGBM在19个基线中具有最低的总体MAE和QS。联合方法的MAE为37.783，较LightGBM增加0.681，即1.84%；QS为24.392，增加0.201，即0.83%。其地区平均RMSE由277.064降至275.397，改善约0.60%。相对XGBoost，联合方法的MAE与QS分别降低约0.39%和4.80%。这些结果表明，联合方法接近强树基线，但尚未获得全面优势。')
fig('comparison_mae.png','19个基线与联合方法的宏平均MAE')
p('LightGBM+QRA的RMSE为269.466，低于联合方法，但其MAE和QS更高，说明不同指标反映了不同误差权衡。深度网络与基础模型在本文固定配置下也未自动获得更优表现；尤其是PriceFM本地移植及基础模型受限输入，使该比较不能直接推广为各模型原始论文能力的排序。')
h('5.2 不同地区的预测表现')
rows=[]
for region in ['NSW1','QLD1','TAS1']:
 for m in ['baseline_lightgbm','baseline_xgboost','Full_joint']:
  r=reg[(reg.model==m)&(reg.region==region)].iloc[0];rows.append([region,label(m),fmt(r.mae_mean,r.mae_std),fmt(r.quantile_score_proxy_mean,r.quantile_score_proxy_std),f'{r.coverage_90_mean*100:.2f}%'])
table('三个地区的核心模型比较',['地区','模型','MAE↓','QS↓','90%覆盖率'],rows,[1.4,4.1,3.5,3.5,2.3],9)
fig('regional_comparison.png','不同地区的MAE与QS对比')
p('在NSW1，联合方法将MAE由44.333降至44.036，QS由31.318降至30.854，两个指标均改善。在QLD1，MAE由38.213升至40.746，QS由24.919升至25.956，成为总体退化的主要来源。TAS1的MAE小幅降低而QS略有增加。QLD1的MAE种子标准差为1.461，也显示出较强的训练波动。')
h('5.3 概率区间的覆盖与宽度')
fig('coverage.png','预测区间覆盖率与名义水平的比较')
p('联合方法的80%和90%区间覆盖率分别为68.56%和82.70%，低于LightGBM的72.50%和85.10%；平均90%区间宽度由141.35缩至118.08。区间变窄未伴随更可靠的覆盖，表明位置校正与尺度控制仍存在失配。QLD1的80%覆盖率从78.08%降至66.73%，是最明显的地区性问题。')
p('CP扩大了区间并提高覆盖率，但其QS较基础树更差，说明覆盖偏高不等同于概率预测更准确。对联合方法同样需要同时审查区间宽度、覆盖与分位损失，避免将某一单项指标的改善解释为整体分布质量提升。')
h('5.4 组件关闭干预分析')
p('相对同一检查点的Anchor输出，完整历史分支将MAE从37.908降至37.783、QS从24.467降至24.392，表现出一定补偿能力。然而，关闭事件项后的MAE为37.754，低于完整模型；关闭时间移动后的QS为24.384，也略低于完整模型的24.392。')
p('上述实验属于固定参数下的关闭干预，并非独立重训消融，因此尚不足以确定某个模块在重新优化后是否无效。但这些结果已表明，当前配置下没有形成事件分支与时间移动的稳定总体净收益。进一步验证需要控制状态编码、损失和训练预算，重新训练无历史、仅趋势及事件无移动等对照。')
h('5.5 电价尖峰与常态误差分析')
p('为分析总体误差背后的状态差异，进一步对全部225份冻结预测进行价格分层评分。预测来源哈希、标签和起点均保持一致，逐格总体MAE复算误差小于10⁻⁸。表5-3与表5-4采用按地区等权后对seed平均的结果，极端价格组是高价组的子集。')
tail=json.loads((P/'tail_report_payload.json').read_text())
for t in tail['tables']:table('价格分层'+('MAE' if 'MAE' in t['title'] else 'QS')+'比较',t['headers'],t['rows'],[4.6,2.8,2.8,2.8,2.8],9)
p('在高价y>300单元上，联合方法的MAE从LightGBM的988.399降至967.603，改善2.10%；在极端y>1000单元上，从3472.494降至3444.268，改善0.81%。但常态0—300区间的MAE从22.822升至23.948，退化4.93%。负价MAE在地区等权平均下改善6.79%，这一结果不意味着每个地区均改善，QLD1负价仍然退化。')
fig('tail_tradeoff.png','联合方法相对LightGBM的价格分层MAE变化')
p('高价与极端区间的QS分别从808.368升至815.337、从2983.550升至3034.190，说明尖峰点误差收益尚未转化为上尾概率预测收益。这与共同平移分位数的结构限制相一致，但仍需通过独立结构对照验证其具体原因。')
p('联合方法也不是尖峰切片上的最佳模型。LightGBM+QRA在高价组、DLinear在极端组具有更低MAE，PatchTST两个尖峰切片均优于联合方法，但总体误差更大。尖峰样本量亦较有限：NSW1、QLD1和TAS1实际超过1000的唯一交付小时分别为85、81和18个。滚动窗口产生的重复评分单元不应作为独立事件，后续统计分析需考虑连续事件段与时间相关性。')
h('5.6 校正稳健性与状态约束实验')
for text in method['sections'][5]['paragraphs']:p(text)
t=method['sections'][5]['table'];table('2021—2022年开发折的稳健性结果',t['headers'],t['rows'],font=9)
p('开发折结果与2023—2024年总体比较采用不同时间段，前者用于分析方案稳定性，不能与后者按绝对误差直接排序。校正限幅通过种子均值门槛而未通过全部单seed门槛，表明其存在保守化收益，但尚未建立跨随机初始化的一致改善。')
h('5.7 结果讨论与有效性限制')
p('总体结果支持一个明确判断：成熟历史能够在部分极端价格上补偿树模型的误差，但在常态价格、地区迁移和概率校准上的代价仍然显著。当前主要问题并非缺少模型复杂度，而是校正何时有可靠证据、应调整位置还是尺度，以及事件修正是否具有独立价值。')
p('实验尚未进行等预算超参数搜索，也未完成预训练重叠和实际首次公开时间的全面审计；相关年份已被多轮研究观察，因而结果属于回顾性分析。三个神经种子不等同于三个独立市场样本，小幅差异亦未经配对时间块显著性检验。后续需要在冻结方案后使用未参与开发的时间段确认外推表现。')

chapter(6,'结论与展望')
h('6.1 阶段研究结论')
p('本研究围绕电价尖峰与概率预测，完成了统一实验流程、19个基线适配及趋势—事件联合校正方法的构建。实验表明，LightGBM仍是总体MAE和QS的最强基线；当前联合方法在部分尖峰点误差和地区上具有收益，但尚不足以抵消常态误差增加与上尾概率预测退化。')
p('方法演进揭示了从历史表征到未来预测之间的关键区别。历史曲线可被清晰分解，并不保证其误差模式能迁移到当前市场。Anchor直接监督、非对称尺度和校正限幅改善了模型约束方式，但其稳定性尚未得到充分证实；固定检查点的干预结果也未支持事件时间移动具有一致增量收益。')
h('6.2 当前研究不足')
p('目前的历史分支对所有分位作共同平移，缺少独立的事件发生概率及条件幅度表示；校正可靠性也尚未与充分的状态支持建立稳定关联。另一方面，尖峰事件数有限、预测窗口重叠及基础模型信息范围差异，使平均结果不足以完整刻画风险。现阶段应将方法定位为尚在验证的研究方案，而非已解决尖峰预测问题的成熟模型。')
h('6.3 后续研究方向')
nexttitles=['误差来源与条件可靠性分析','状态证据约束及统一收缩对照','历史事件增量价值的重训消融','事件概率与条件强度分离建模','跨时期确认与统计检验']
for i,(s,title) in enumerate(zip(method['next_steps'],nexttitles),1):
 h(f'6.3.{i} {title}',3)
 p(s['design'])
 p('实验规模为'+s['matrix'].rstrip('。')+'。'+s['decision'])
h('6.4 总结')
p('下一阶段的核心目标是在保留强基座常态精度的前提下，获得可重复的尖峰风险改善。判断新结构是否有效，应同时考察总体误差、尖峰发生与强度、误报代价及条件覆盖，并通过相同信息范围和训练预算的对照区分数据价值、校正收缩与结构创新。只有在这些证据形成一致支持后，才能进一步提出稳定的模型贡献。')

hd=h('参考文献',1);hd.paragraph_format.page_break_before=True
for i,r in enumerate(refs,1):
 pa=p(f"[{i}] {r['citation']}",indent=False)
 for run in pa.runs:run.font.superscript=False;run.font.size=Pt(10.5)
 pa.paragraph_format.line_spacing=1.15;pa.paragraph_format.space_after=Pt(6);pa.paragraph_format.left_indent=Cm(.7);pa.paragraph_format.first_line_indent=Cm(-.7)
 if r.get('url'):
  from docx.opc.constants import RELATIONSHIP_TYPE as RT
  pa.add_run(' ');link=OxmlElement('w:hyperlink');link.set(qn('r:id'),pa.part.relate_to(r['url'],RT.HYPERLINK,is_external=True));run=OxmlElement('w:r');prop=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'19');prop.append(size);run.append(prop);txt=OxmlElement('w:t');txt.text=r['url'];run.append(txt);link.append(run);pa._p.append(link)
chapter_num='A';tab_num=0
hd=h('附录A  补充实验结果',1);hd.paragraph_format.page_break_before=True
p('表A-1补充旧静态Full与无场结构的同协议结果。它们属于历史结构对照，与当前联合方法的模型身份不同。表A-2提供全部25个输出的价格分层MAE。')
hist=pd.read_csv(D/'historical_controls.csv');table('历史静态结构对照',['模型','MAE','QS','RMSE'],[[r.model,fmt(r.mae_mean,r.mae_std),fmt(r.quantile_score_proxy_mean,r.quantile_score_proxy_std),fmt(r.rmse_pooled_mean)] for _,r in hist.iterrows()],[4.5,4,4,3.3],9)
tdf=pd.read_csv(P/'latest_spike_diagnostics_macro.csv').set_index(['model','band'])
table('全部输出的价格分层MAE',['模型/输出','负价<0','常态0–300','高价>300','极端>1000'],[[label(m)]+[f'{tdf.loc[(m,b),"mae_mean"]:.3f}' for b in ['negative','normal_0_to_300','high_gt300','extreme_gt1000']] for m in main.model],[4.8,2.75,2.75,2.75,2.75],9)
chapter_num='B';tab_num=0
hd=h('附录B  可复现性说明',1);hd.paragraph_format.page_break_before=True
p('总体结果来自冻结的2023—2024年预测文件。结果整理版本为2026年10月24日维护后的统一比较集，当前联合方法版本为asymmetric_balance_20261023；本文另对同一批预测作价格分层复算。数据路径、来源哈希和逐格分数保存在配套材料中，未在报告编制过程中重新训练或基于测试结果调整模型。')
p('基线参考文献按原始出版页和论文记录核对。预训练基础模型区分模型世代、论文来源与实际权重版本，相关本地书目信息的更正记录保存在文献元数据文件中。Informer和iTransformer虽保留历史实现，但未进入本次19个已完成基线的总体比较。')
table('配套复现材料',['材料','用途'],[['main.csv / regional.csv','总体及分地区评分'],['per_seed.csv / verification.json','逐seed结果与来源一致性核验'],['latest_spike_diagnostics_macro.csv','本次价格分层宏平均评分'],['latest_spike_diagnostics_provenance.json','分层复算来源、哈希与样本核验'],['baselines.json / method.json','方法来源与版本记录']],[6,9.8],10)
out=P/'电价尖峰预测研究_阶段报告_论文章节版.docx';doc.save(out)
(P/'academic_manifest.json').write_text(json.dumps({'file':str(out),'chapters':6,'baseline_count':len(baselines['models']),'reference_count':len(refs),'figures':len(doc.inline_shapes),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False,indent=2))
print(out)
